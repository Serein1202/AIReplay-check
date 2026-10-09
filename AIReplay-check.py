#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Replay 核查一体化脚本（图形界面版）
=================================================================
一次运行完成「后台抓取 → CSV 导出 → 筛选/排序/去重 → xlsx 输出」全流程，
带可视化操作界面（customtkinter），替代原有的原生弹窗交互。

界面用法:
  1. 双击/运行脚本打开窗口，抓取日期默认「昨天」，可点「昨天/今天/前天」或手动输入
  2. 点击「开始抓取并处理」，日志区实时滚动显示抓取与处理进度
  3. 完成后弹出结果卡片，统计筛选、去重行数，可一键打开输出目录

命令行用法:
  python AIReplay-check.py 20260924         # 静默模式（自动化 / 定时任务）
  python AIReplay-check.py --date 20260924  # 同上
  python AIReplay-check.py --selfcheck      # 界面自检（启动后自动退出）
  python AIReplay-check.py --check-update   # 只检测更新并打印结果（不启动界面）
  python AIReplay-check.py --version        # 打印当前版本号

登录态说明:
  获取顺序（优先级由高到低）:
    1. 环境变量 AMIS_TOKEN
    2. 本地 Token 缓存（脚本 / exe 同目录下的 .amis_token.json，未过期时直接复用）
    3. 本机 Chrome / Edge 中 amis.ssjj.cn 的有效 Token（AMIS-USER-TOKEN）
    4. 以上均不可用（未找到或已过期）时，自动弹出登录框，手动登录换取新 Token 后继续抓取
  本地 Token 缓存:
    - 登录成功（含自动弹框登录）后统一写入缓存，内容仅 token / account / expire / saved_at
    - 缓存命中且未过期 → 直接使用并在日志中提示来源与有效期，不再要求重新登录
    - 使用缓存 / 浏览器 Token 抓取时若被服务端判定失效（HTTP 401/403 或返回鉴权、
      登录失效类错误码 / 消息），自动删除缓存文件并弹出登录框换取新 Token，
      登录成功后覆盖写入缓存并继续本次抓取，无需用户重跑
    - 缓存文件损坏或解析失败时按无缓存处理，并删除该文件
    - 密码仅在内存中使用，绝不写入缓存、日志或代码
  登录框:
    - 框内展示登录网址 http://amis.ssjj.cn/，含账号、密码（掩码显示）与「LDAP 登录」
      复选项（默认不勾选，使用 LDAP 账号时勾选）
    - 调用 POST http://amis.ssjj.cn/auth/signIn（Header X-PROJECT-ID: 1，
      体 {"account","password","ldap"}）换取 Token
    - 登录失败会提示服务端返回的 msg 并允许重试（最多 3 次）；全部失败或取消则中止本次抓取
    - 有界面（customtkinter）时用深色对话框（CTkToplevel）；无界面时退回原生 tkinter 对话框
    - 账号密码仅在本机内存中使用，不写入磁盘、日志或代码；脚本内不含任何硬编码凭据

在线更新:
  - 打开程序后自动检测 GitHub 最新 Release；发现更高版本时弹窗提示，可一键下载并重启到新版本
  - 新版本下载到 exe 同目录，文件名与 Release 资产一致（AIReplay-<版本号>.exe）；
    重启后由新版本自动清理旧版本文件，目录里不会堆积历史 exe
  - 网络不可达（如无法访问 api.github.com）时静默跳过，不影响正常使用
  - 关闭自动检测: 设置环境变量 AIREPLAY_NO_UPDATE=1，或把 UPDATE_AUTO_CHECK 改为 False
  - 版本号来源: 打包时由 GitHub Actions 注入 APP_VERSION（见 build-release.yml）；
    以源码方式运行时显示占位版本，且「立即更新」会改为打开 Releases 下载页

依赖:
  pip install pandas openpyxl sentry-sdk customtkinter

打包 exe:
  pyinstaller --noconsole --onefile --collect-all customtkinter --name "AIReplay" AIReplay-check.py

自动构建与发布:
  向 main 分支推送任何 .py 改动 → GitHub Actions 自动打包 exe 并发布到仓库 Releases，
  发布文件命名为 AIReplay-<版本号>.exe（如 AIReplay-v1.0.1.exe）
  配置见 .github/workflows/build-release.yml，版本号默认 v1.0.<构建号>；
  也可以打标签指定版本: git tag v1.2.0 && git push origin v1.2.0
"""

import base64
import csv
import json
import os
import queue
import re
import shutil
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.parse
import urllib.request
import webbrowser
from datetime import date, datetime, timedelta

# ====================== 可配置项 ======================
# 后台接口
HOST = "http://amis.ssjj.cn"
API_PATH = "/gameConfigs/listAiHackCheckReport"
PROJECT_ID = "28"
STATUS_DONE = 3          # 3 = 已核查
PER_PAGE = 1000          # 单页拉取条数（实测 1000 可一次取全）
HTTP_TIMEOUT = 30        # 单次请求超时（秒）

# 登录接口（自动获取 Token 失败 / Token 过期时的兜底：手动登录换取新 Token）
LOGIN_PAGE = HOST + "/"          # 登录页展示地址：http://amis.ssjj.cn/
LOGIN_API = HOST + "/auth/signIn"
LOGIN_PROJECT_ID = "1"           # 登录接口专用项目号，固定 "1"（用 PROJECT_ID=28 会返回 400）
LOGIN_MAX_RETRY = 3              # 手动登录最大重试次数

# 本地 Token 缓存：保存于脚本 / exe 同目录，避免每次抓取都重新登录
TOKEN_CACHE_NAME = ".amis_token.json"

# 输出命名（YYYYMMDD 由输入日期填充）
CSV_NAME_TPL = "replay已核查-{day}.csv"
XLSX_NAME_TPL = "replay核查-{day}.xlsx"

# 目标文件已存在时，是否先重命名保留旧文件（True=防覆盖，False=直接覆盖）
PROTECT_EXISTING = True

# 筛选条件：列名含以下关键字（不区分大小写），精确值匹配
FILTER_COLUMNS = {
    "封禁状态": "未封禁",
    "备注":     "无异常",
}
# 排序：列名含以下关键字（不区分大小写），默认降序
SORT_COLUMN_KEYWORD = "分数"
# 去重依据：列名含以下关键字（不区分大小写），保留第一次出现的行
DEDUP_COLUMN_KEYWORD = "角色ID"
# 想设为"文本"的列：按表头关键字匹配（不区分大小写），命中即转换
TEXT_COLUMN_KEYWORDS = ["id", "房间"]
# 想设为"日期时间"的列：按表头关键字匹配（不区分大小写），命中即转换
DATETIME_COLUMN_KEYWORDS = ["开始时间", "结束时间", "上报时间"]
# 也可直接写死列字母（优先级最高），例如 ["A", "C"]
TEXT_COLUMN_LETTERS = []

# CSV 表头（与页面表格列一一对应）
CSV_HEADERS = [
    "#", "处理状态", "角色名", "角色ID", "无端ID", "房间号", "分数",
    "开始时间", "战斗耗时", "上报时间", "处理人", "处理时间", "封禁状态", "备注",
]
# 处理状态码 → 文案
STATUS_TEXT = {3: "已核查"}

# ====================== Sentry 监控配置 ======================
SENTRY_DSN = "https://1164181e93905c40acb94ed3c27679c1@o4511234196570112.ingest.us.sentry.io/4512146535940096"
SENTRY_TRACES_SAMPLE_RATE = 1.0
# ============================================================

APP_TITLE = "Replay 核查工具"
FONT_UI = "Microsoft YaHei UI"

# ====================== 版本与在线更新 ======================
# 当前版本号：仓库源码里保留占位值，打包时由 GitHub Actions 注入该次发布的真实版本号
# （见 .github/workflows/build-release.yml 的「注入版本号」步骤）
APP_VERSION = "0.0.0"
APP_VERSION_PLACEHOLDER = "0.0.0"

# 在线更新：查询该仓库的最新 Release，发现更高版本时提示并可自动下载、重启到新版本
GITHUB_REPO = "Serein1202/AIReplay-check"
UPDATE_API = f"https://api.github.com/repos/{GITHUB_REPO}/releases/latest"
UPDATE_TIMEOUT = 15                # 检测 / 下载超时（秒）
UPDATE_ASSET_PREFIX = "AIReplay-"  # Release 中 exe 资产的文件名前缀
UPDATE_AUTO_CHECK = True           # 启动时自动检测更新（环境变量 AIREPLAY_NO_UPDATE=1 可临时关闭）
UPDATE_MARKER = ".aireplay_update.json"  # 记录上次在线更新后待清理的旧文件

# ---- 以下是内部实现，一般无需修改 ----

try:
    if SENTRY_DSN:
        import sentry_sdk
        sentry_sdk.init(
            dsn=SENTRY_DSN,
            traces_sample_rate=SENTRY_TRACES_SAMPLE_RATE,
            environment="production",
            release=os.path.basename(__file__),
        )
except Exception as _e:  # Sentry 不可用不影响主流程
    try:
        print(f"[警告] Sentry 初始化失败（不影响运行）: {_e}", flush=True)
    except Exception:
        pass


def report_exception(exc, **tags):
    """把异常上报到 Sentry（未配置 DSN 或未安装 SDK 时静默跳过）"""
    try:
        if not SENTRY_DSN:
            return
        import sentry_sdk
        if tags:
            sentry_sdk.set_tags(tags)
        sentry_sdk.capture_exception(exc)
    except Exception:
        pass


# ---------------------------------------------------------------- 基础工具
def script_dir():
    """脚本所在目录（PyInstaller 打包成 exe 时取 exe 所在目录）"""
    if getattr(sys, "frozen", False):
        return os.path.dirname(os.path.abspath(sys.executable))
    return os.path.dirname(os.path.abspath(__file__))


def _console(msg):
    """控制台输出；exe 无控制台（--noconsole）时静默忽略"""
    try:
        print(msg, flush=True)
    except Exception:
        pass


def _exit_now(code=0):
    """立即结束进程，跳过解释器收尾阶段，保证关窗口后进程马上消失。

    onefile 打包的程序在退出时，解释器收尾（atexit 回调、非守护线程等待、
    大模块析构）加上杀毒软件对解包目录的扫描，会让进程在窗口关闭后仍残留
    很久。这里只做必要的收尾（上报 flush、标准输出 flush），然后直接结束。
    """
    try:
        n = int(code or 0)
    except Exception:
        n = 0
    # Sentry 若已加载，给它一次短暂的机会把事件发出去
    _sentry = sys.modules.get("sentry_sdk")
    if _sentry is not None:
        try:
            _sentry.flush(timeout=1.5)
        except Exception:
            pass
    for _stream in (sys.stdout, sys.stderr):
        try:
            _stream.flush()
        except Exception:
            pass
    os._exit(n)


# 界面日志接收器：GUI 启动时被替换为向界面推送的回调，其余场景为 None
_LOG_SINK = None


def log(msg):
    """统一日志出口：控制台 + 界面日志区"""
    _console(msg)
    sink = _LOG_SINK
    if sink is not None:
        try:
            sink(f"{datetime.now():%H:%M:%S}  {msg}")
        except Exception:
            pass


def popup(title, text, kind="info"):
    """原生弹窗提示（仅在无 GUI 依赖的降级路径中使用）"""
    try:
        import tkinter as tk
        from tkinter import messagebox
        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        getattr(messagebox, kind)(title, text)
        root.destroy()
    except Exception:
        log(f"[{title}] {text}")


def ask_date(default_day):
    """原生弹窗输入抓取日期，返回 YYYYMMDD 字符串；取消返回 None"""
    import tkinter as tk
    from tkinter import simpledialog, messagebox
    root = tk.Tk()
    root.withdraw()
    root.attributes("-topmost", True)
    while True:
        val = simpledialog.askstring(
            "抓取日期",
            "请输入抓取日期（格式：20260924）",
            initialvalue=default_day,
            parent=root,
        )
        if val is None:
            root.destroy()
            return None
        cleaned = re.sub(r"[^0-9]", "", val)
        try:
            datetime.strptime(cleaned, "%Y%m%d")
            root.destroy()
            return cleaned
        except ValueError:
            messagebox.showwarning(
                "格式错误", "日期格式不正确，请按 20260924 的格式重新输入。", parent=root
            )


def backup_if_exists(path):
    """目标文件已存在时先重命名保留，避免覆盖既有产物"""
    if not (PROTECT_EXISTING and os.path.exists(path)):
        return
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base, ext = os.path.splitext(path)
    new_path = f"{base}（旧-{stamp}）{ext}"
    try:
        os.replace(path, new_path)
        log(f"[提示] 目标文件已存在，已重命名保留旧文件: {new_path}")
    except Exception as e:
        log(f"[警告] 旧文件重命名失败（将继续写入）: {e}")


def open_folder(path):
    """在资源管理器中打开目录。

    [重要] Windows 的 ShellExecute 对「不带扩展名」的路径会先按 PATHEXT 规则
    逐项尝试补扩展名（.exe/.com/.bat/...），全部落空后才按目录处理。
    所以当输出目录 ...\\replay核查\\replay核查 的同级存在 replay核查.exe 时，
    os.startfile(path) 会被解析成「启动 replay核查.exe」——现象就是
    「点打开输出目录却又新运行了一个 exe」，而不是打开目录。
    修复: Windows 下改用 explorer 打开；并以追加尾分隔符作为兜底。
    """
    import subprocess

    path = os.path.normpath(path)
    os.makedirs(path, exist_ok=True)

    if hasattr(os, "startfile"):  # Windows
        try:
            subprocess.Popen(["explorer", path])
        except OSError:
            # 兜底: 尾分隔符强制按目录处理，避免被同名 .exe 抢走
            os.startfile(path + os.sep)
        log(f"[提示] 已在资源管理器中打开目录: {path}")
    else:
        subprocess.Popen(["xdg-open", path])
        log(f"[提示] 已打开目录: {path}")
    return path


# ==================================================================
#                      版本与在线更新（GitHub Releases）
# ==================================================================
def current_exe_path():
    """当前可执行文件路径；以源码方式（python xxx.py）运行时返回 None"""
    if getattr(sys, "frozen", False):
        return os.path.abspath(sys.executable)
    return None


def current_app_version():
    """当前版本号（不含前导 v）。

    优先级：
      1. APP_VERSION —— 打包时由 CI 注入，代表这个 exe 的真实版本
      2. exe 文件名里的版本（AIReplay-v1.0.1.exe → 1.0.1），用于本地打包 / 手工改名场景
    """
    v = (APP_VERSION or "").strip().lstrip("vV")
    if v and v != APP_VERSION_PLACEHOLDER:
        return v
    if getattr(sys, "frozen", False):
        m = re.search(r"AIReplay[-_]?v?(\d+(?:\.\d+)*)",
                      os.path.basename(sys.executable), re.I)
        if m:
            return m.group(1)
    return v or APP_VERSION_PLACEHOLDER


def _version_key(value):
    """把版本号转成可比较的元组：1.0.10 > 1.0.9"""
    nums = re.findall(r"\d+", str(value))
    return tuple(int(n) for n in nums[:6]) if nums else (0,)


def version_is_newer(latest, current):
    """latest 是否比 current 更新"""
    a, b = _version_key(latest), _version_key(current)
    n = max(len(a), len(b))
    return a + (0,) * (n - len(a)) > b + (0,) * (n - len(b))


def check_for_update():
    """查询 GitHub 最新 Release。

    返回 (status, payload)：
      ("update", info)  有可用更新。info = {version, tag, name, url, asset_name, size, notes}
      ("latest", info)  已是最新版本
      ("error",  msg)   检测失败（网络不可达 / 尚无 Release 等），msg 为可读说明
    """
    req = urllib.request.Request(
        UPDATE_API,
        headers={
            "Accept": "application/vnd.github+json",
            "User-Agent": f"AIReplay-check/{current_app_version()}",
        },
    )
    try:
        with urllib.request.urlopen(req, timeout=UPDATE_TIMEOUT) as resp:
            data = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        if e.code == 404:
            return ("error", "仓库还没有发布任何 Release")
        if e.code == 403:
            return ("error", "GitHub 接口访问受限（可能是请求过于频繁），请稍后再试")
        return ("error", f"GitHub 返回 HTTP {e.code}")
    except Exception as e:
        return ("error", f"无法访问 GitHub（{type(e).__name__}: {e}）")

    tag = (data.get("tag_name") or "").strip()
    assets = data.get("assets") or []

    asset = None
    for a in assets:  # 优先取 AIReplay-*.exe
        nm = (a.get("name") or "")
        if nm.lower().endswith(".exe") and nm.lower().startswith(UPDATE_ASSET_PREFIX.lower()):
            asset = a
            break
    if asset is None:  # 退而求其次：任意 .exe
        for a in assets:
            if (a.get("name") or "").lower().endswith(".exe"):
                asset = a
                break
    if asset is None:
        return ("error", f"Release {tag or '(无标签)'} 中没有可用的 exe 文件")

    info = {
        "version": tag.lstrip("vV"),
        "tag": tag,
        "name": data.get("name") or tag,
        "url": asset.get("browser_download_url") or "",
        "asset_name": asset.get("name") or "",
        "size": int(asset.get("size") or 0),
        "notes": data.get("body") or "",
    }
    if version_is_newer(info["version"], current_app_version()):
        return ("update", info)
    return ("latest", info)


def download_update(info, on_progress=None):
    """下载新版本 exe 到程序所在目录，返回新文件路径。

    先写入 .part 再改名，避免下载中断留下的半成品被误当成可用程序。
    """
    url = info.get("url") or ""
    name = info.get("asset_name") or ""
    if not url or not name:
        raise RuntimeError("下载地址或文件名缺失")

    dest = os.path.join(script_dir(), os.path.basename(name))
    tmp = dest + ".part"
    try:
        os.remove(tmp)
    except OSError:
        pass

    req = urllib.request.Request(
        url, headers={"User-Agent": f"AIReplay-check/{current_app_version()}"}
    )
    got = 0
    try:
        with urllib.request.urlopen(req, timeout=UPDATE_TIMEOUT) as resp, open(tmp, "wb") as f:
            total = int(resp.headers.get("Content-Length") or info.get("size") or 0)
            while True:
                chunk = resp.read(256 * 1024)
                if not chunk:
                    break
                f.write(chunk)
                got += len(chunk)
                if on_progress:
                    try:
                        on_progress(got, total)
                    except Exception:
                        pass
    except Exception:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise

    # 完整性校验：正常产物约 35 MB，明显偏小视为失败
    if got < 1024 * 1024:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise RuntimeError(f"下载内容不完整（仅 {got} 字节）")
    if total and got != total:
        try:
            os.remove(tmp)
        except OSError:
            pass
        raise RuntimeError(f"下载不完整：{got} / {total} 字节")

    os.replace(tmp, dest)
    log(f"[更新] 新版本已下载完成: {dest}")
    return dest


def launch_detached(path):
    """以独立进程启动指定程序（不随当前进程退出而被回收）"""
    import subprocess
    try:
        if hasattr(os, "startfile"):
            os.startfile(path)   # noqa: S606 —— Windows：等同双击运行
            return
    except OSError:
        pass
    subprocess.Popen([path], cwd=os.path.dirname(path) or None)


def prepare_update_cleanup(old_exe):
    """记录旧版本文件，交由新版本启动时清理（避免目录里堆积历史 exe）"""
    try:
        with open(os.path.join(script_dir(), UPDATE_MARKER), "w", encoding="utf-8") as f:
            json.dump({"remove": [old_exe]}, f)
    except Exception:
        pass


def cleanup_after_update():
    """启动时清理上次在线更新遗留的旧版本文件。

    新版本刚启动时旧 exe 可能仍被占用，因此带重试；仍失败则保留记录，下次启动再试。
    安全约束：只删程序所在目录下的文件，且绝不删当前正在运行的程序。
    """
    marker = os.path.join(script_dir(), UPDATE_MARKER)
    if not os.path.exists(marker):
        return
    try:
        with open(marker, encoding="utf-8") as f:
            data = json.load(f)
    except Exception:
        data = {}

    folder = os.path.abspath(script_dir())
    current = current_exe_path()
    pending = []

    for path in data.get("remove") or []:
        target = os.path.abspath(str(path))
        if os.path.dirname(target) != folder:      # 安全：只处理同目录文件
            continue
        if current and target == current:          # 安全：绝不删正在运行的程序
            continue
        for attempt in range(10):
            try:
                os.remove(target)
                break
            except FileNotFoundError:
                break
            except Exception:
                if attempt == 9:
                    pending.append(str(path))
                else:
                    time.sleep(1)

    if pending:
        try:
            with open(marker, "w", encoding="utf-8") as f:
                json.dump({"remove": pending}, f)
        except Exception:
            pass
        return
    try:
        os.remove(marker)
    except OSError:
        pass


# ---------------------------------------------------------------- 登录 Token
JWT_RE = re.compile(rb"eyJ[A-Za-z0-9_\-+/=]+\.eyJ[A-Za-z0-9_\-+/=]+\.[0-9a-fA-F]{64}")
BROWSER_ROOTS = [
    os.path.join(os.environ.get("LOCALAPPDATA", ""), "Google", "Chrome", "User Data"),
    os.path.join(os.environ.get("LOCALAPPDATA", ""), "Microsoft", "Edge", "User Data"),
]


def parse_token_info(token):
    """解析 JWT 中间段，取 account / expire 等字段"""
    try:
        payload = token.split(".")[1]
        payload += "=" * (-len(payload) % 4)
        return json.loads(base64.urlsafe_b64decode(payload.encode("ascii")).decode("utf-8"))
    except Exception:
        return {}


def scan_leveldb_dir(leveldb_dir, tmp_root):
    """复制 LevelDB 文件后扫描，返回 [(expire, token, info), ...]"""
    found = []
    tmp_dir = os.path.join(tmp_root, "ls_" + str(abs(hash(leveldb_dir)) % 10 ** 8))
    os.makedirs(tmp_dir, exist_ok=True)
    for name in os.listdir(leveldb_dir):
        src = os.path.join(leveldb_dir, name)
        if not os.path.isfile(src) or name == "LOCK":
            continue
        dst = os.path.join(tmp_dir, name)
        try:
            shutil.copy2(src, dst)
            blob = open(dst, "rb").read()
        except Exception:
            continue
        for m in JWT_RE.finditer(blob):
            ctx = blob[max(0, m.start() - 600): m.start()]
            if b"amis.ssjj.cn" not in ctx and b"AMIS-USER-TOKEN" not in ctx:
                continue
            token = m.group(0).decode("ascii")
            info = parse_token_info(token)
            if info.get("account"):
                found.append((int(info.get("expire") or 0), token, info))
    return found


# ---------------------------------------------------------------- 本地 Token 缓存
# 缓存文件位于脚本 / exe 同目录，仅保存 token / account / expire / saved_at，
# 绝不保存密码；Token 失效或被服务端判定失效时自动删除，下次重新登录后再写入。
def token_cache_path():
    """本地 Token 缓存文件路径（脚本 / exe 同目录下的 .amis_token.json）"""
    return os.path.join(script_dir(), TOKEN_CACHE_NAME)


def _fmt_expire(expire):
    """把 expire 时间戳格式化为可读文本；无效时返回「未知」"""
    try:
        expire = int(expire or 0)
    except (TypeError, ValueError):
        return "未知"
    if not expire:
        return "未知"
    try:
        return datetime.fromtimestamp(expire).strftime("%Y-%m-%d %H:%M:%S")
    except Exception:
        return "未知"


def cache_is_fresh(info):
    """缓存 Token 是否仍在有效期内

    expire 缺失（无法解析出有效期）时视为可用：真失效会被服务端判定，
    届时自动删缓存并弹框重新登录，不会造成不可恢复的死循环。
    """
    try:
        expire = int((info or {}).get("expire") or 0)
    except (TypeError, ValueError):
        return False
    if not expire:
        return True
    return expire > datetime.now().timestamp()


def load_token_cache():
    """读取本地 Token 缓存；返回 (token, info)；无缓存/损坏时返回 None

    文件不存在：返回 None（按无缓存处理，不打印告警）。
    文件损坏或解析失败：删除该文件并按无缓存处理。
    """
    path = token_cache_path()
    if not os.path.isfile(path):
        return None
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        if not isinstance(data, dict):
            raise ValueError("缓存内容不是 JSON 对象")
        token = str(data.get("token") or "").strip()
        if not token:
            raise ValueError("缓存中缺少 token 字段")
        info = parse_token_info(token)
        account = str(data.get("account") or "").strip()
        if account:
            info["account"] = account
        try:
            expire = int(data.get("expire") or 0)
        except (TypeError, ValueError):
            expire = 0
        if expire:
            info["expire"] = expire
        return token, info
    except Exception as e:
        log(f"[警告] 本地 Token 缓存损坏或无法解析，已按无缓存处理并删除: {e}")
        remove_token_cache("缓存损坏")
        return None


def save_token_cache(token, info):
    """写入本地 Token 缓存（仅 token / account / expire / saved_at，不含密码）"""
    if not token:
        return False
    info = info or {}
    try:
        expire = int(info.get("expire") or 0)
    except (TypeError, ValueError):
        expire = 0
    payload = {
        "token": token,
        "account": str(info.get("account") or ""),
        "expire": expire,
        "saved_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }
    path = token_cache_path()
    try:
        with open(path, "w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        log(f"[缓存] 已写入本地 Token 缓存（账号 {payload['account'] or '未知'}，"
            f"有效期至 {_fmt_expire(expire)}）: {path}")
        return True
    except Exception as e:
        log(f"[警告] 本地 Token 缓存写入失败（不影响本次抓取）: {e}")
        return False


def remove_token_cache(reason=""):
    """删除本地 Token 缓存文件（Token 失效 / 缓存损坏时调用）"""
    path = token_cache_path()
    if not os.path.isfile(path):
        return False
    try:
        os.remove(path)
        note = f"（{reason}）" if reason else ""
        log(f"[缓存] 已删除本地 Token 缓存{note}: {path}")
        return True
    except Exception as e:
        log(f"[警告] 本地 Token 缓存删除失败: {e}")
        return False


class TokenInvalidError(RuntimeError):
    """服务端判定登录态失效（HTTP 401/403 或返回鉴权 / 登录失效类错误码、消息）"""


# 出现以下消息关键字组合时，视为登录态失效
_AUTH_FAIL_PHRASES = (
    "未登录", "请先登录", "请重新登录", "重新登录", "登录已失效", "登录失效",
    "登录过期", "登录状态失效", "鉴权失败", "鉴权错误", "认证失败", "未授权",
    "无效的token", "token无效", "token 无效", "token已过期", "token 已过期",
    "令牌无效", "令牌已失效",
)


def looks_like_auth_failure(code=None, msg=""):
    """判断接口返回是否属于「Token / 登录态失效」类错误"""
    try:
        code = int(code)
    except (TypeError, ValueError):
        code = None
    if code in (401, 403):
        return True
    text = str(msg or "").strip()
    if not text:
        return False
    low = text.lower().replace(" ", "")
    if any(p.replace(" ", "") in low for p in _AUTH_FAIL_PHRASES):
        return True
    has_token_word = ("token" in low) or ("令牌" in text)
    has_auth_word = any(k in text for k in ("登录", "登陆", "失效", "过期", "鉴权", "认证", "授权"))
    return has_token_word and has_auth_word


def find_amis_token():
    """按优先级获取 amis.ssjj.cn 登录 Token

    顺序：环境变量 AMIS_TOKEN > 本地缓存 .amis_token.json（未过期）>
    浏览器 Chrome / Edge Local Storage（取有效期最长者）。
    返回 (token, info, source)；均不可用时返回 (None, {}, "")，由上层弹登录框。
    """
    env_token = os.environ.get("AMIS_TOKEN", "").strip()
    if env_token:
        info = parse_token_info(env_token)
        return env_token, info, "环境变量 AMIS_TOKEN"

    cached = load_token_cache()
    if cached:
        token, info = cached
        if cache_is_fresh(info):
            return token, info, "本地缓存"
        # 缓存已过期：删除后继续走后续获取路径
        remove_token_cache("本地缓存已过期")

    tmp_root = tempfile.mkdtemp(prefix="replay_token_")
    candidates = []
    try:
        for root in BROWSER_ROOTS:
            if not os.path.isdir(root):
                continue
            for profile in os.listdir(root):
                leveldb = os.path.join(root, profile, "Local Storage", "leveldb")
                if os.path.isdir(leveldb):
                    candidates.extend(scan_leveldb_dir(leveldb, tmp_root))
    finally:
        shutil.rmtree(tmp_root, ignore_errors=True)

    if not candidates:
        return None, {}, ""
    candidates.sort(key=lambda x: x[0], reverse=True)
    _, token, info = candidates[0]
    return token, info, "浏览器 Local Storage"


# ---------------------------------------------------------------- 登录兜底（Token 失效时手动登录）
# 登录框 UI 挂载点：GUI 启动时注册「主线程弹框」回调；未注册时回退原生 tkinter 对话框
_LOGIN_UI = None


def login_by_password(account, password, ldap=True):
    """调用后台登录接口换取新 Token（与界面完全解耦）

    成功返回 (token, info)；失败抛出 RuntimeError，消息为服务端返回的 msg。
    账号密码仅在本函数调用期间存在于内存中，不做任何持久化。
    """
    body = json.dumps({
        "account": account,
        "password": password,
        "ldap": bool(ldap),
    }).encode("utf-8")
    req = urllib.request.Request(
        LOGIN_API,
        data=body,
        headers={
            "Content-Type": "application/json",
            # 注意: /auth/signIn 的 X-PROJECT-ID 必须固定为 "1"（实测用 PROJECT_ID=28
            # 或不带该头均返回 HTTP 400 Bad Request，用 "1" 才返回正常 JSON）。
            # 数据抓取接口的 X-PROJECT-ID 仍为 PROJECT_ID(28)，不要混用。
            "X-PROJECT-ID": LOGIN_PROJECT_ID,
            "Accept": "application/json, text/plain, */*",
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                          "(KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36",
            "Referer": LOGIN_PAGE,
        },
        method="POST",
    )
    try:
        with urllib.request.urlopen(req, timeout=HTTP_TIMEOUT) as resp:
            payload = json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as e:
        raise RuntimeError(f"登录接口返回 HTTP {e.code}（{e.reason}）")
    except Exception as e:
        raise RuntimeError(f"登录请求发送失败：{e}")

    if payload.get("code") != 0:
        raise RuntimeError(payload.get("msg") or f"登录失败（code={payload.get('code')}）")

    token = str(((payload.get("data") or {}).get("token")) or "").strip()
    if not token:
        raise RuntimeError("登录成功但响应中未包含 Token（data.token 为空）")

    info = parse_token_info(token)
    if not info.get("account"):
        info["account"] = account
    return token, info


def ask_credentials(tip="", error="", default_account=""):
    """弹出登录框收集凭据；返回 (account, password, ldap)，用户取消返回 None

    UI 与登录逻辑解耦：GUI 模式下走 _LOGIN_UI 注册的深色对话框，
    其余场景回退到原生 tkinter 对话框。
    """
    ui = _LOGIN_UI
    if ui is not None:
        return ui(tip, error, default_account)
    return _ask_credentials_tk(tip, error, default_account)


def _ask_credentials_tk(tip="", error="", default_account=""):
    """无 customtkinter 时的原生 tkinter 登录对话框（兜底）"""
    import tkinter as tk

    result = {"value": None}
    root = tk.Tk()
    root.title("登录 amis.ssjj.cn")
    root.resizable(False, False)
    root.attributes("-topmost", True)

    tk.Label(root, text="登录 amis.ssjj.cn",
             font=("Microsoft YaHei UI", 13, "bold")).grid(
        row=0, column=0, columnspan=2, sticky="w", padx=14, pady=(14, 2))
    tk.Label(root, text=tip or "自动获取 Token 失败，请手动登录后继续抓取。",
             fg="#666666", justify="left", wraplength=360).grid(
        row=1, column=0, columnspan=2, sticky="w", padx=14, pady=(0, 8))

    tk.Label(root, text="登录网址").grid(row=2, column=0, sticky="e", padx=14, pady=6)
    url_row = tk.Frame(root)
    url_row.grid(row=2, column=1, sticky="w", padx=14, pady=6)
    tk.Label(url_row, text=LOGIN_PAGE, fg="#2F6BFF").pack(side="left")
    tk.Button(url_row, text="打开",
              command=lambda: webbrowser.open(LOGIN_PAGE)).pack(side="left", padx=(8, 0))

    tk.Label(root, text="账号").grid(row=3, column=0, sticky="e", padx=14, pady=6)
    var_account = tk.StringVar(value=default_account)
    ent_account = tk.Entry(root, textvariable=var_account, width=30)
    ent_account.grid(row=3, column=1, sticky="w", padx=14, pady=6)

    tk.Label(root, text="密码").grid(row=4, column=0, sticky="e", padx=14, pady=6)
    var_password = tk.StringVar()
    ent_password = tk.Entry(root, textvariable=var_password, width=30, show="●")
    ent_password.grid(row=4, column=1, sticky="w", padx=14, pady=6)

    var_ldap = tk.BooleanVar(value=False)
    tk.Checkbutton(root, text="LDAP 登录（本后台使用 LDAP 账号时请勾选）", variable=var_ldap).grid(
        row=5, column=0, columnspan=2, sticky="w", padx=14, pady=(2, 2))

    lbl_msg = tk.Label(root, text=error, fg="#EF4444", justify="left", wraplength=360)
    lbl_msg.grid(row=6, column=0, columnspan=2, sticky="w", padx=14, pady=(2, 4))

    def _submit():
        account = var_account.get().strip()
        password = var_password.get()
        if not account or not password:
            lbl_msg.configure(text="请输入账号和密码。")
            return
        # 凭据仅放入内存结果，不做任何持久化
        result["value"] = (account, password, bool(var_ldap.get()))
        root.destroy()

    def _cancel():
        result["value"] = None
        root.destroy()

    btns = tk.Frame(root)
    btns.grid(row=7, column=0, columnspan=2, sticky="e", padx=14, pady=(6, 14))
    tk.Button(btns, text="登录", width=12, command=_submit).pack(side="left", padx=(0, 8))
    tk.Button(btns, text="取消", width=8, command=_cancel).pack(side="left")

    root.protocol("WM_DELETE_WINDOW", _cancel)
    root.bind("<Return>", lambda _e: _submit())
    root.bind("<Escape>", lambda _e: _cancel())
    ent_account.focus_set()

    root.update_idletasks()
    x = (root.winfo_screenwidth() - root.winfo_width()) // 2
    y = (root.winfo_screenheight() - root.winfo_height()) // 3
    root.geometry(f"+{max(0, x)}+{max(0, y)}")
    root.mainloop()
    return result["value"]


def obtain_token_by_login(reason=""):
    """自动获取 Token 失败或已过期时，弹框手动登录换取 Token（最多重试 3 次）

    成功返回 (token, info)；用户取消或重试耗尽则抛 RuntimeError。
    账号密码仅内存使用，不写入磁盘 / 日志，脚本内也不含任何硬编码凭据。
    """
    default_account = ""
    error = ""
    for attempt in range(1, LOGIN_MAX_RETRY + 1):
        tip = reason if attempt == 1 else ""
        creds = ask_credentials(tip=tip, error=error, default_account=default_account)
        if creds is None:
            raise RuntimeError(
                "已取消登录：未取得 amis.ssjj.cn 的有效 Token，本次抓取中止。\n"
                "可先在 Chrome / Edge 登录 amis.ssjj.cn，或设置环境变量 AMIS_TOKEN 后再试。"
            )
        account, password, ldap = creds
        default_account = account
        try:
            token, info = login_by_password(account, password, ldap)
        except Exception as e:
            error = str(e)
            log(f"[登录失败] 第 {attempt}/{LOGIN_MAX_RETRY} 次：{error}")
            if attempt < LOGIN_MAX_RETRY:
                continue
            raise RuntimeError(
                f"连续 {LOGIN_MAX_RETRY} 次登录失败，本次抓取中止。\n"
                f"最后一次错误：{error}\n"
                "请确认账号 / 密码 / LDAP 选项是否正确，或优先在浏览器登录 amis.ssjj.cn。"
            )
        log(f"[登录成功] 已获取 amis.ssjj.cn 登录 Token"
            f"（第 {attempt} 次尝试，LDAP={'是' if ldap else '否'}）。")
        # 登录成功统一写入本地缓存（仅 Token 与账号，密码不落盘）
        save_token_cache(token, info)
        return token, info


# ---------------------------------------------------------------- 数据抓取
def _fetch_pages(headers, start_at, end_at):
    """按分页抓取全部记录；Token 失效时抛 TokenInvalidError，由上层换 Token 后重试"""
    items, page, total = [], 1, None
    while True:
        qs = urllib.parse.urlencode({
            "page": page,
            "perPage": PER_PAGE,
            "date_type": 2,
            "start_at": start_at,
            "end_at": end_at,
            "status": STATUS_DONE,
        })
        url = f"{HOST}{API_PATH}?{qs}"
        req = urllib.request.Request(url, headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=HTTP_TIMEOUT) as resp:
                payload = json.loads(resp.read().decode("utf-8"))
        except urllib.error.HTTPError as e:
            if e.code in (401, 403):
                raise TokenInvalidError(
                    f"接口返回 HTTP {e.code}（{e.reason}），登录态已失效"
                )
            raise RuntimeError(f"接口请求失败：HTTP {e.code}（{e.reason}）")
        except urllib.error.URLError as e:
            raise RuntimeError(f"接口请求失败：{e.reason}")
        if payload.get("code") != 0:
            msg = payload.get("msg")
            if looks_like_auth_failure(payload.get("code"), msg):
                raise TokenInvalidError(
                    f"接口返回鉴权失败（code={payload.get('code')}，msg={msg}），登录态已失效"
                )
            raise RuntimeError(f"接口返回异常: {msg or payload}")
        data = payload.get("data") or {}
        batch = data.get("items") or []
        total = data.get("total") if total is None else total
        items.extend(batch)
        log(f"[抓取] 第 {page} 页 {len(batch)} 条，累计 {len(items)}/{total} 条")
        if not batch or (total is not None and len(items) >= int(total)):
            break
        page += 1
    return items


def _log_token_source(source, info):
    """日志提示本次 Token 的来源与有效期"""
    expire = int((info or {}).get("expire") or 0)
    detail = f"，有效期至 {_fmt_expire(expire)}" if expire else ""
    log(f"[登录态] 已自动获取有效 Token（来源：{source or '未知'}{detail}）。")


def fetch_day(day):
    """抓取指定日期 00:00:00 ~ 次日 00:00:00 的已核查记录

    Token 获取顺序：环境变量 AMIS_TOKEN > 本地缓存 .amis_token.json（未过期）>
    浏览器 Local Storage 有效 Token > 弹出登录框（手动输入账号密码换取新 Token）。
    使用缓存 / 浏览器 Token 抓取时若被服务端判定失效，会自动删除缓存并弹出登录框，
    登录成功后覆盖写入缓存并继续本次抓取，无需用户重跑。
    """
    token, info, source = find_amis_token()
    reason = ""
    if not token:
        reason = ("未能自动获取 amis.ssjj.cn 登录 Token"
                  "（未设置环境变量 AMIS_TOKEN，本地缓存与浏览器中也没有找到可用 Token）")
    else:
        expire = int(info.get("expire") or 0)
        if expire and expire < datetime.now().timestamp():
            reason = "自动获取到的 amis.ssjj.cn 登录 Token 已过期"

    if reason:
        log(f"[提示] {reason}，请在登录窗口中输入账号密码换取新 Token。")
        token, info = obtain_token_by_login(reason)
        source = "登录框"
    else:
        _log_token_source(source, info)

    account = info.get("account") or "xhdong"

    start_at = f"{day[0:4]}-{day[4:6]}-{day[6:8]} 00:00:00"
    end_at = (datetime.strptime(day, "%Y%m%d") + timedelta(days=1)).strftime("%Y-%m-%d 00:00:00")

    items = []
    auth_retry = 0
    while True:
        headers = {
            "X-AMIS-TOKEN": token,
            "X-AMIS-ACCOUNT": account,
            "X-PROJECT-ID": PROJECT_ID,
            "Accept": "application/json, text/plain, */*",
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                          "(KHTML, like Gecko) Chrome/140.0.0.0 Safari/537.36",
            "Referer": HOST + "/",
        }
        try:
            items = _fetch_pages(headers, start_at, end_at)
            break
        except TokenInvalidError as e:
            log(f"[登录态失效] {e}")
            # 服务端判定失效：删除本地缓存并弹框换新 Token
            remove_token_cache("服务端判定登录态失效")
            if auth_retry >= LOGIN_MAX_RETRY:
                raise RuntimeError(
                    f"登录 Token 连续 {LOGIN_MAX_RETRY} 次被服务端判定失效，本次抓取中止。\n"
                    f"最后一次错误：{e}"
                )
            auth_retry += 1
            log("[提示] 缓存 / 自动获取的 Token 已失效，请在登录窗口中重新登录以换取新 Token。")
            token, info = obtain_token_by_login("自动获取的登录 Token 已失效，请重新登录")
            account = info.get("account") or account
            # 登录成功已覆盖写入缓存，继续本次抓取
    return items, start_at, end_at


def item_to_row(it):
    return [
        it.get("id", ""),
        STATUS_TEXT.get(it.get("status"), str(it.get("status", ""))),
        it.get("role_name", ""),
        it.get("role_id", ""),
        it.get("wd_id", ""),
        it.get("room_id", ""),
        it.get("score", ""),
        it.get("start_date", ""),
        it.get("duration", ""),
        it.get("report_at", ""),
        it.get("operator", ""),
        it.get("deal_at", ""),
        "已封禁" if it.get("forbid") else "未封禁",
        it.get("remarks", ""),
    ]


def write_csv(rows, out_path):
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    with open(out_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, quoting=csv.QUOTE_ALL)
        writer.writerow(CSV_HEADERS)
        writer.writerows(rows)
    return out_path


# ---------------------------------------------------------------- 文档处理（逻辑内联自 replay核查.py）
def _column_error(col_keyword, step):
    msg = f"错误：【{step}】未找到包含「{col_keyword}」的列，请检查表头名称是否正确。"
    log(msg)
    raise RuntimeError(msg)


def detect_csv_encoding(path, sample=8192):
    for enc in ("utf-8-sig", "gb18030", "utf-8", "gbk"):
        try:
            with open(path, "r", encoding=enc) as f:
                f.read(sample)
            return enc
        except (UnicodeDecodeError, UnicodeError):
            continue
    return "latin-1"


def find_col_by_keyword(headers, keyword):
    kw = keyword.lower()
    for h in headers:
        if kw in str(h).strip().lower():
            return str(h).strip()
    return None


def resolve_text_columns(headers, ncols):
    from openpyxl.utils import column_index_from_string
    idx_set = set()
    for i, h in enumerate(headers):
        h_str = str(h).strip().lower()
        if any(kw.lower() in h_str for kw in TEXT_COLUMN_KEYWORDS):
            idx_set.add(i)
    for letter in TEXT_COLUMN_LETTERS:
        ci = column_index_from_string(letter) - 1
        if 0 <= ci < ncols:
            idx_set.add(ci)
    return idx_set


def resolve_datetime_columns(headers):
    idx_set = set()
    for i, h in enumerate(headers):
        h_str = str(h).strip().lower()
        if any(kw.lower() in h_str for kw in DATETIME_COLUMN_KEYWORDS):
            idx_set.add(i)
    return idx_set


def auto_fit_column_width(ws, padding=4, min_width=8, max_width=50):
    from openpyxl.utils import get_column_letter
    for col_cells in ws.iter_cols():
        col_letter = get_column_letter(col_cells[0].column)
        max_len = 0
        for cell in col_cells:
            val = cell.value
            text = "" if val is None else str(val)
            fullwidth = len(re.findall(r"[\u4e00-\u9fff\uff00-\uffef]", text))
            ascii_len = len(text) - fullwidth
            max_len = max(max_len, fullwidth * 2 + ascii_len)
        ws.column_dimensions[col_letter].width = min(max_width, max(min_width, max_len + padding))


def process_csv(path, out_folder, day):
    """CSV → 筛选/排序/去重 → xlsx，返回统计信息与输出路径"""
    import pandas as pd
    from openpyxl.utils import get_column_letter

    enc = detect_csv_encoding(path)
    header_df = pd.read_csv(path, nrows=0, encoding=enc)
    headers = list(header_df.columns)
    text_cols = resolve_text_columns(headers, len(headers))
    datetime_cols = resolve_datetime_columns(headers)
    if not text_cols:
        _column_error("'id' 或 房间", "文本列")

    converters = {headers[i]: str for i in text_cols}
    for i in datetime_cols:
        converters[headers[i]] = pd.to_datetime

    df = pd.read_csv(path, encoding=enc, converters=converters, keep_default_na=False)

    # 筛选
    rows_before = len(df)
    for col_key, val_key in FILTER_COLUMNS.items():
        col = find_col_by_keyword(df.columns, col_key)
        if col is None:
            _column_error(col_key, "筛选")
        df = df[df[col].astype(str).str.strip() == str(val_key).strip()]
    rows_after_filter = len(df)

    # 排序
    sort_col = find_col_by_keyword(df.columns, SORT_COLUMN_KEYWORD)
    if sort_col is None:
        _column_error(SORT_COLUMN_KEYWORD, "排序")
    df[sort_col] = pd.to_numeric(df[sort_col], errors="coerce")
    df = df.sort_values(by=sort_col, ascending=False, na_position="last")

    # 去重
    dedup_col = find_col_by_keyword(df.columns, DEDUP_COLUMN_KEYWORD)
    rows_before_dedup = len(df)
    if dedup_col is None:
        _column_error(DEDUP_COLUMN_KEYWORD, "去重")
    df = df.drop_duplicates(subset=[dedup_col], keep="first")
    rows_after_dedup = len(df)

    df = df.reset_index(drop=True)

    # 输出路径
    os.makedirs(out_folder, exist_ok=True)
    out_path = os.path.join(out_folder, XLSX_NAME_TPL.format(day=day))
    backup_if_exists(out_path)

    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Sheet1")
        ws = writer.sheets["Sheet1"]
        for ci in text_cols:
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=ci + 1)
                if cell.value is None or (
                    isinstance(cell.value, str) and cell.value.strip().lower() == "nan"
                ):
                    cell.value = ""
                else:
                    cell.value = str(cell.value)
                cell.number_format = "@"
        for ci in datetime_cols:
            for r in range(2, ws.max_row + 1):
                cell = ws.cell(row=r, column=ci + 1)
                if cell.value is None:
                    cell.value = ""
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
        from openpyxl.styles import Alignment
        center = Alignment(horizontal="center", vertical="center")
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.alignment = center
        auto_fit_column_width(ws)

    return {
        "encoding": enc,
        "rows_before": rows_before,
        "rows_after_filter": rows_after_filter,
        "rows_before_dedup": rows_before_dedup,
        "rows_after_dedup": rows_after_dedup,
        "out_path": out_path,
    }


# ---------------------------------------------------------------- 主流程（界面与命令行共用）
class PipelineError(RuntimeError):
    """带阶段标识的流程异常（用于 Sentry 打标签与界面提示）"""

    def __init__(self, stage, message):
        super().__init__(message)
        self.stage = stage


def run_pipeline(day):
    """抓取 → CSV → 处理 → xlsx，返回结果字典（异常转 PipelineError）"""
    base = script_dir()
    csv_folder = os.path.join(base, "csv")
    xlsx_folder = os.path.join(base, "replay核查")

    log(f"===== 目标日期: {day} =====")

    # 1. 抓取
    try:
        items, start_at, end_at = fetch_day(day)
    except Exception as e:
        report_exception(e, stage="fetch", day=str(day))
        log(f"[错误] 抓取失败: {e}")
        raise PipelineError("fetch", str(e))
    if not items:
        log(f"[结束] {start_at} ~ {end_at} 没有「已核查」记录。")
        return {"status": "empty", "day": day, "start_at": start_at, "end_at": end_at}

    # 2. 导出 CSV
    try:
        csv_path = write_csv(
            [item_to_row(it) for it in items],
            os.path.join(csv_folder, CSV_NAME_TPL.format(day=day)),
        )
    except Exception as e:
        report_exception(e, stage="csv", day=str(day))
        log(f"[错误] CSV 导出失败: {e}")
        raise PipelineError("csv", str(e))
    log(f"[抓取完成] 共 {len(items)} 条 → {csv_path}")

    # 3. 处理文档
    try:
        stat = process_csv(csv_path, xlsx_folder, day)
    except Exception as e:
        report_exception(e, stage="process", day=str(day))
        log(f"[错误] 处理失败: {e}")
        raise PipelineError("process", str(e))

    log(f"[文档处理完成] {stat['out_path']}")
    return {
        "status": "ok",
        "day": day,
        "count": len(items),
        "start_at": start_at,
        "end_at": end_at,
        "csv_path": csv_path,
        "stat": stat,
    }


# ==================================================================
#                          图形界面（customtkinter）
# ==================================================================
try:
    import customtkinter as ctk
except Exception:  # 未安装 customtkinter 时自动退回原生弹窗模式
    ctk = None

# 配色（浅色, 深色）
CLR = {
    "win":      ("#F3F5FA", "#0F1216"),
    "card":     ("#FFFFFF", "#181B21"),
    "line":     ("#E4E8F0", "#242830"),
    "text":     ("#111827", "#EEF1F6"),
    "sub":      ("#667085", "#8B94A6"),
    "accent":   ("#2F6BFF", "#3B82F6"),
    "accent_h": ("#1E54E0", "#2C6FE0"),
    "ghost":    ("#ECF0F7", "#222731"),
    "ghost_h":  ("#DFE6F2", "#2B313D"),
    "log_bg":   ("#FAFBFD", "#0B0D11"),
}
OK_CLR, WARN_CLR, ERR_CLR, RUN_CLR = "#22C55E", "#F59E0B", "#EF4444", "#3B82F6"

if ctk is not None:

    class ResultWindow(ctk.CTkToplevel):
        """结果卡片：展示统计信息与输出文件，可一键打开目录"""

        def __init__(self, master, res):
            super().__init__(master)
            self.res = res
            ok = res.get("status") == "ok"
            stat = res.get("stat") or {}

            self.title("核查完成" if ok else "无数据")
            self.configure(fg_color=CLR["win"])
            self.resizable(False, False)
            self.geometry("640x420")

            f_title = ctk.CTkFont(family=FONT_UI, size=19, weight="bold")
            f_key = ctk.CTkFont(family=FONT_UI, size=12)
            f_val = ctk.CTkFont(family=FONT_UI, size=13, weight="bold")
            f_mono = ctk.CTkFont(family="Consolas", size=11)

            head = ctk.CTkFrame(self, fg_color="transparent")
            head.pack(fill="x", padx=24, pady=(22, 6))
            ctk.CTkLabel(
                head, text=("✓  核查完成" if ok else "!  该日期无数据"),
                font=f_title, text_color=(OK_CLR if ok else WARN_CLR),
            ).pack(side="left")

            body = ctk.CTkFrame(self, fg_color=CLR["card"], corner_radius=14,
                                border_width=1, border_color=CLR["line"])
            body.pack(fill="both", expand=True, padx=24, pady=(10, 0))
            body.grid_columnconfigure(1, weight=1)

            if ok:
                rows = [
                    ("抓取日期", res.get("day", "")),
                    ("抓取范围", f"{res.get('start_at','')}  ~  {res.get('end_at','')}"),
                    ("抓取条数", f"{res.get('count', 0)} 条"),
                    ("筛选结果", f"{stat.get('rows_before', 0)}  →  {stat.get('rows_after_filter', 0)} 行"
                                 f"（未封禁 + 无异常）"),
                    ("去重结果", f"{stat.get('rows_before_dedup', 0)}  →  {stat.get('rows_after_dedup', 0)} 行"
                                 f"（按角色ID）"),
                    ("排序方式", "按分数降序"),
                ]
            else:
                rows = [
                    ("抓取日期", res.get("day", "")),
                    ("抓取范围", f"{res.get('start_at','')}  ~  {res.get('end_at','')}"),
                    ("状态", "该时间段内没有「已核查」记录"),
                ]
            for i, (k, v) in enumerate(rows):
                ctk.CTkLabel(body, text=k, font=f_key, text_color=CLR["sub"]).grid(
                    row=i, column=0, sticky="w", padx=(20, 16), pady=(14 if i == 0 else 7, 0))
                ctk.CTkLabel(body, text=v, font=f_val, text_color=CLR["text"], anchor="w",
                             justify="left").grid(
                    row=i, column=1, sticky="w", padx=(0, 20), pady=(14 if i == 0 else 7, 0))

            out_path = stat.get("out_path") or res.get("csv_path") or ""
            if out_path:
                ctk.CTkLabel(body, text="输出文件", font=f_key, text_color=CLR["sub"]).grid(
                    row=len(rows), column=0, columnspan=2, sticky="w", padx=20, pady=(18, 4))
                ent = ctk.CTkEntry(body, height=34, font=f_mono, corner_radius=9,
                                   fg_color=CLR["log_bg"], border_color=CLR["line"],
                                   text_color=CLR["text"])
                ent.grid(row=len(rows) + 1, column=0, columnspan=2, sticky="ew", padx=20, pady=(0, 18))
                ent.insert(0, out_path)
                ent.configure(state="readonly")

            foot = ctk.CTkFrame(self, fg_color="transparent")
            foot.pack(fill="x", padx=24, pady=18)
            if out_path:
                ctk.CTkButton(
                    foot, text="打开输出目录", width=150, height=40, corner_radius=10,
                    font=ctk.CTkFont(family=FONT_UI, size=13, weight="bold"),
                    fg_color=CLR["accent"], hover_color=CLR["accent_h"], text_color="#FFFFFF",
                    command=lambda p=os.path.dirname(out_path): self._open(p),
                ).pack(side="left")
            ctk.CTkButton(
                foot, text="关闭", width=100, height=40, corner_radius=10,
                font=ctk.CTkFont(family=FONT_UI, size=13),
                fg_color=CLR["ghost"], hover_color=CLR["ghost_h"], text_color=CLR["text"],
                command=self.destroy,
            ).pack(side="right")

            self.transient(master)
            self.after(120, self._place)
            self.lift()
            try:
                self.grab_set()
            except Exception:
                pass

        def _place(self):
            try:
                self.update_idletasks()
                m = self.master
                x = m.winfo_rootx() + (m.winfo_width() - self.winfo_width()) // 2
                y = m.winfo_rooty() + (m.winfo_height() - self.winfo_height()) // 3
                self.geometry(f"+{max(0, x)}+{max(0, y)}")
            except Exception:
                pass

        def _open(self, path):
            try:
                open_folder(path)
            except Exception as e:
                log(f"[警告] 打开目录失败: {e}")

    class LoginDialog(ctk.CTkToplevel):
        """登录对话框（深色）：账号 + 密码（掩码）+ LDAP 复选（默认不勾选）

        只负责收集界面输入，返回 (account, password, ldap)；用户取消返回 None。
        实际的登录请求由 login_by_password() 完成，二者互不耦合。
        """

        def __init__(self, master, tip="", error="", default_account=""):
            super().__init__(master)
            self._result = None

            self.title("登录 amis.ssjj.cn")
            self.configure(fg_color=CLR["win"])
            self.resizable(False, False)
            self.geometry("470x420")

            f_h1 = ctk.CTkFont(family=FONT_UI, size=18, weight="bold")
            f_key = ctk.CTkFont(family=FONT_UI, size=12)
            f_sub = ctk.CTkFont(family=FONT_UI, size=11)
            f_entry = ctk.CTkFont(family="Consolas", size=14)
            f_btn = ctk.CTkFont(family=FONT_UI, size=13, weight="bold")

            head = ctk.CTkFrame(self, fg_color="transparent")
            head.pack(fill="x", padx=24, pady=(20, 4))
            ctk.CTkLabel(head, text="登录 amis.ssjj.cn", font=f_h1,
                         text_color=CLR["text"]).pack(anchor="w")
            ctk.CTkLabel(head, text=tip or "自动获取 Token 失败，请手动登录后继续抓取。",
                         font=f_sub, text_color=CLR["sub"], wraplength=410,
                         justify="left").pack(anchor="w", pady=(4, 0))

            card = ctk.CTkFrame(self, fg_color=CLR["card"], corner_radius=14,
                                border_width=1, border_color=CLR["line"])
            card.pack(fill="both", expand=True, padx=24, pady=(12, 0))
            card.grid_columnconfigure(1, weight=1)

            # 登录网址（明确展示）
            ctk.CTkLabel(card, text="登录网址", font=f_key,
                         text_color=CLR["sub"]).grid(row=0, column=0, sticky="w",
                                                     padx=(20, 12), pady=(16, 4))
            url_row = ctk.CTkFrame(card, fg_color="transparent")
            url_row.grid(row=0, column=1, sticky="ew", padx=(0, 16), pady=(16, 4))
            url_row.grid_columnconfigure(0, weight=1)
            ctk.CTkLabel(url_row, text=LOGIN_PAGE, font=f_entry,
                         text_color=CLR["accent"]).grid(row=0, column=0, sticky="w")
            ctk.CTkButton(url_row, text="打开", width=56, height=28, corner_radius=8,
                          font=f_sub, fg_color=CLR["ghost"], hover_color=CLR["ghost_h"],
                          text_color=CLR["text"],
                          command=self._open_login_page).grid(row=0, column=1, sticky="e")

            # 账号
            ctk.CTkLabel(card, text="账号", font=f_key,
                         text_color=CLR["sub"]).grid(row=1, column=0, sticky="w",
                                                     padx=(20, 12), pady=(10, 0))
            self.ent_account = ctk.CTkEntry(card, height=36, font=f_entry, corner_radius=9,
                                            fg_color=CLR["log_bg"], border_color=CLR["line"],
                                            text_color=CLR["text"], placeholder_text="请输入账号")
            self.ent_account.grid(row=1, column=1, sticky="ew", padx=(0, 16), pady=(10, 0))

            # 密码（掩码显示）
            ctk.CTkLabel(card, text="密码", font=f_key,
                         text_color=CLR["sub"]).grid(row=2, column=0, sticky="w",
                                                     padx=(20, 12), pady=(10, 0))
            self.ent_password = ctk.CTkEntry(card, height=36, font=f_entry, corner_radius=9,
                                             fg_color=CLR["log_bg"], border_color=CLR["line"],
                                             text_color=CLR["text"], show="●",
                                             placeholder_text="请输入密码")
            self.ent_password.grid(row=2, column=1, sticky="ew", padx=(0, 16), pady=(10, 0))

            # LDAP 登录（默认不勾选）
            self.var_ldap = ctk.BooleanVar(value=False)
            ctk.CTkCheckBox(
                card, text="LDAP 登录（本后台使用 LDAP 账号时请勾选）", variable=self.var_ldap,
                font=f_sub, text_color=CLR["text"], checkbox_width=18, checkbox_height=18,
                corner_radius=5, fg_color=CLR["accent"], hover_color=CLR["accent_h"],
                border_color=CLR["line"], checkmark_color="#FFFFFF",
            ).grid(row=3, column=0, columnspan=2, sticky="w", padx=20, pady=(14, 4))

            # 服务端错误提示（重试时回显上一次失败原因）
            self.lbl_error = ctk.CTkLabel(card, text=error or "", font=f_sub,
                                          text_color=ERR_CLR, wraplength=390, justify="left")
            self.lbl_error.grid(row=4, column=0, columnspan=2, sticky="w",
                                padx=20, pady=(6, 16))

            foot = ctk.CTkFrame(self, fg_color="transparent")
            foot.pack(fill="x", padx=24, pady=16)
            ctk.CTkButton(foot, text="登录", width=140, height=40, corner_radius=10,
                          font=f_btn, fg_color=CLR["accent"], hover_color=CLR["accent_h"],
                          text_color="#FFFFFF",
                          command=self._submit).pack(side="left")
            ctk.CTkButton(foot, text="取消", width=100, height=40, corner_radius=10,
                          font=f_btn, fg_color=CLR["ghost"], hover_color=CLR["ghost_h"],
                          text_color=CLR["text"],
                          command=self._cancel).pack(side="right")

            self.protocol("WM_DELETE_WINDOW", self._cancel)
            self.bind("<Return>", lambda _e: self._submit())
            self.bind("<Escape>", lambda _e: self._cancel())
            if default_account:
                self.ent_account.insert(0, default_account)
            self.transient(master)
            self.after(120, self._place)
            self.lift()
            try:
                self.grab_set()
            except Exception:
                pass
            (self.ent_password if default_account else self.ent_account).focus_set()

        def _place(self):
            try:
                self.update_idletasks()
                m = self.master
                x = m.winfo_rootx() + (m.winfo_width() - self.winfo_width()) // 2
                y = m.winfo_rooty() + (m.winfo_height() - self.winfo_height()) // 3
                self.geometry(f"+{max(0, x)}+{max(0, y)}")
            except Exception:
                pass

        def _open_login_page(self):
            try:
                webbrowser.open(LOGIN_PAGE)
            except Exception as e:
                log(f"[警告] 打开登录页失败: {e}")

        def _submit(self):
            account = self.ent_account.get().strip()
            password = self.ent_password.get()
            if not account:
                self.lbl_error.configure(text="请输入账号。")
                return
            if not password:
                self.lbl_error.configure(text="请输入密码。")
                return
            # 凭据仅放入内存结果，不做任何持久化
            self._result = (account, password, bool(self.var_ldap.get()))
            self.destroy()

        def _cancel(self):
            self._result = None
            self.destroy()

        def wait(self):
            """阻塞至对话框关闭，返回 (account, password, ldap) 或 None（取消）"""
            self.wait_window(self)
            return self._result

    class UpdateDialog(ctk.CTkToplevel):
        """发现新版本：展示版本信息并支持在线更新"""

        def __init__(self, master, info):
            super().__init__(master)
            self.app = master
            self.info = info
            self._busy = False

            self.title("发现新版本")
            self.configure(fg_color=CLR["win"])
            self.resizable(False, False)
            self.geometry("620x480")

            f_title = ctk.CTkFont(family=FONT_UI, size=19, weight="bold")
            f_key = ctk.CTkFont(family=FONT_UI, size=12)
            f_val = ctk.CTkFont(family=FONT_UI, size=13, weight="bold")
            f_body = ctk.CTkFont(family=FONT_UI, size=12)
            f_mono = ctk.CTkFont(family="Consolas", size=11)

            head = ctk.CTkFrame(self, fg_color="transparent")
            head.pack(fill="x", padx=24, pady=(22, 6))
            ctk.CTkLabel(head, text="↑  发现新版本", font=f_title,
                         text_color=CLR["accent"]).pack(side="left")

            body = ctk.CTkFrame(self, fg_color=CLR["card"], corner_radius=14,
                                border_width=1, border_color=CLR["line"])
            body.pack(fill="both", expand=True, padx=24, pady=(10, 0))
            body.grid_columnconfigure(1, weight=1)

            size = info.get("size") or 0
            rows = [
                ("当前版本", f"v{current_app_version()}"),
                ("最新版本", f"v{info.get('version', '')}"),
                ("文件大小", f"{size / 1048576:.1f} MB" if size else "未知"),
            ]
            for i, (k, v) in enumerate(rows):
                ctk.CTkLabel(body, text=k, font=f_key, text_color=CLR["sub"]).grid(
                    row=i, column=0, sticky="w", padx=(20, 16), pady=(14 if i == 0 else 7, 0))
                ctk.CTkLabel(body, text=v, font=f_val, text_color=CLR["text"],
                             anchor="w").grid(row=i, column=1, sticky="w",
                                              padx=(0, 20), pady=(14 if i == 0 else 7, 0))

            notes = (info.get("notes") or "").strip()
            ctk.CTkLabel(body, text="更新说明", font=f_key, text_color=CLR["sub"]).grid(
                row=len(rows), column=0, columnspan=2, sticky="w", padx=20, pady=(18, 4))
            txt = ctk.CTkTextbox(body, height=126, corner_radius=10, font=f_mono, wrap="word",
                                 fg_color=CLR["log_bg"], border_width=1,
                                 border_color=CLR["line"], text_color=CLR["text"])
            txt.grid(row=len(rows) + 1, column=0, columnspan=2, sticky="nsew",
                     padx=20, pady=(0, 16))
            txt.insert("1.0", notes[:4000] if notes else "（该版本未提供更新说明）")
            txt.configure(state="disabled")
            body.grid_rowconfigure(len(rows) + 1, weight=1)

            self.lbl_msg = ctk.CTkLabel(self, text="", font=f_body, text_color=CLR["sub"])
            self.lbl_msg.pack(fill="x", padx=26, pady=(10, 0))
            self.progress = ctk.CTkProgressBar(self, height=6, corner_radius=3,
                                               progress_color=CLR["accent"],
                                               fg_color=CLR["ghost"])
            self.progress.set(0)
            self.progress.pack(fill="x", padx=26, pady=(6, 0))

            foot = ctk.CTkFrame(self, fg_color="transparent")
            foot.pack(fill="x", padx=24, pady=18)
            self.btn_go = ctk.CTkButton(
                foot, text="立即更新", width=150, height=40, corner_radius=10,
                font=ctk.CTkFont(family=FONT_UI, size=13, weight="bold"),
                fg_color=CLR["accent"], hover_color=CLR["accent_h"], text_color="#FFFFFF",
                command=self._start_update)
            self.btn_go.pack(side="left")
            self.btn_later = ctk.CTkButton(
                foot, text="稍后", width=100, height=40, corner_radius=10, font=f_body,
                fg_color=CLR["ghost"], hover_color=CLR["ghost_h"], text_color=CLR["text"],
                command=self._close)
            self.btn_later.pack(side="right")

            self.transient(master)
            self.after(120, self._place)
            self.lift()
            try:
                self.grab_set()
            except Exception:
                pass
            self.protocol("WM_DELETE_WINDOW", self._close)

        def _place(self):
            try:
                self.update_idletasks()
                m = self.master
                x = m.winfo_rootx() + (m.winfo_width() - self.winfo_width()) // 2
                y = m.winfo_rooty() + (m.winfo_height() - self.winfo_height()) // 3
                self.geometry(f"+{max(0, x)}+{max(0, y)}")
            except Exception:
                pass

        def _close(self):
            if self._busy:
                return          # 下载中不允许关闭，避免留下半成品
            self.destroy()

        # ---- 下载与替换 ----
        def _start_update(self):
            if self._busy:
                return
            if not current_exe_path():
                # 源码运行：不做自我替换，改为打开下载页
                self._busy = True
                self.btn_go.configure(state="disabled", text="已打开下载页")
                self.lbl_msg.configure(text="当前为源码运行，请在浏览器中手动下载新版本。")
                try:
                    webbrowser.open(f"https://github.com/{GITHUB_REPO}/releases/latest")
                except Exception:
                    pass
                return
            self._busy = True
            self.btn_go.configure(state="disabled", text="正在下载…")
            self.btn_later.configure(state="disabled")
            self.lbl_msg.configure(text="正在下载新版本，请勿关闭窗口…")
            threading.Thread(target=self._download_worker, daemon=True).start()

        def _download_worker(self):
            try:
                path = download_update(
                    self.info,
                    on_progress=lambda got, total: self.app._q.put(
                        ("update_progress", (got, total))),
                )
            except Exception as e:
                self.app._q.put(("update_error", f"{type(e).__name__}: {e}"))
                return
            self.app._q.put(("update_ready", path))

        def _on_progress(self, got, total):
            try:
                if total:
                    self.progress.set(min(1.0, got / total))
                    self.lbl_msg.configure(
                        text=f"正在下载… {got / 1048576:.1f} / {total / 1048576:.1f} MB")
                else:
                    self.lbl_msg.configure(text=f"正在下载… {got / 1048576:.1f} MB")
            except Exception:
                pass

        def _on_error(self, msg):
            self._busy = False
            try:
                self.btn_go.configure(state="normal", text="重试")
                self.btn_later.configure(state="normal")
                self.lbl_msg.configure(text=f"更新失败：{msg}")
            except Exception:
                pass
            try:
                import tkinter.messagebox as mb
                mb.showerror("更新失败",
                             f"下载新版本失败：\n\n{msg}\n\n"
                             f"可前往 https://github.com/{GITHUB_REPO}/releases 手动下载。")
            except Exception:
                pass

        def _on_ready(self, path):
            try:
                self.progress.set(1)
                self.lbl_msg.configure(text="下载完成，正在重启到新版本…")
                old = current_exe_path()
                if old and os.path.abspath(old) != os.path.abspath(path):
                    prepare_update_cleanup(old)
                launch_detached(path)
            except Exception as e:
                self._busy = False
                try:
                    self.btn_go.configure(state="normal", text="重试")
                    self.lbl_msg.configure(text=f"启动新版本失败：{e}，请手动运行 {path}")
                except Exception:
                    pass
                return
            log(f"[更新] 已启动新版本，当前程序即将退出: {path}")
            self.app.after(700, self.app._quit_for_update)

    class ReplayApp(ctk.CTk):
        """主窗口"""

        def __init__(self, selfcheck=False):
            global _LOG_SINK, _LOGIN_UI
            super().__init__()

            self._q = queue.Queue()
            self._running = False
            self._update_dialog = None
            self._day_default = (date.today() - timedelta(days=1)).strftime("%Y%m%d")

            self.title(f"{APP_TITLE}  v{current_app_version()}")
            self.configure(fg_color=CLR["win"])
            self.geometry(self._center_geometry(900, 720))
            self.minsize(840, 640)

            self.f_h1 = ctk.CTkFont(family=FONT_UI, size=23, weight="bold")
            self.f_sub = ctk.CTkFont(family=FONT_UI, size=12)
            self.f_sec = ctk.CTkFont(family=FONT_UI, size=13, weight="bold")
            self.f_btn = ctk.CTkFont(family=FONT_UI, size=14, weight="bold")
            self.f_body = ctk.CTkFont(family=FONT_UI, size=12)
            self.f_entry = ctk.CTkFont(family="Consolas", size=15)
            self.f_log = ctk.CTkFont(family="Consolas", size=12)
            self.f_status = ctk.CTkFont(family=FONT_UI, size=12, weight="bold")

            self.grid_columnconfigure(0, weight=1)
            self.grid_rowconfigure(3, weight=1)
            self._build_header()
            self._build_settings()
            self._build_actions()
            self._build_log()
            self._build_footer()

            self.protocol("WM_DELETE_WINDOW", self._on_close)
            _LOG_SINK = lambda text: self._q.put(("log", text))
            _LOGIN_UI = self._ask_login_threadsafe  # 注册登录框入口（抓取线程经此回到主线程弹框）
            self._append_log(f"{datetime.now():%H:%M:%S}  界面已就绪，默认抓取日期 {self._day_default}")
            self.after(80, self._drain)
            threading.Thread(target=self._preheat, daemon=True).start()
            # 打开即检测更新（自检模式 / 环境变量关闭时不检测）
            if not selfcheck and UPDATE_AUTO_CHECK and not os.environ.get("AIREPLAY_NO_UPDATE"):
                self.after(1200, lambda: self._check_update(manual=False))

        # ---------------- 布局 ----------------
        def _center_geometry(self, w, h):
            try:
                sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
                return f"{w}x{h}+{max(0, (sw - w) // 2)}+{max(0, (sh - h) // 2 - 30)}"
            except Exception:
                return f"{w}x{h}"

        def _build_header(self):
            bar = ctk.CTkFrame(self, fg_color="transparent")
            bar.grid(row=0, column=0, sticky="ew", padx=24, pady=(20, 12))
            bar.grid_columnconfigure(0, weight=1)

            left = ctk.CTkFrame(bar, fg_color="transparent")
            left.grid(row=0, column=0, sticky="w")
            ctk.CTkLabel(left, text=APP_TITLE, font=self.f_h1,
                         text_color=CLR["text"]).grid(row=0, column=0, sticky="w")
            ctk.CTkLabel(left, text="自动抓取  →  CSV 导出  →  筛选去重  →  Excel 报表",
                         font=self.f_sub, text_color=CLR["sub"]).grid(row=1, column=0, sticky="w", pady=(3, 0))

            right = ctk.CTkFrame(bar, fg_color="transparent")
            right.grid(row=0, column=1, sticky="e")
            self.seg_theme = ctk.CTkSegmentedButton(
                right, values=["深色", "浅色"], width=128, height=30, font=self.f_body,
                command=self._switch_theme, selected_color=CLR["accent"],
                selected_hover_color=CLR["accent_h"], unselected_color=CLR["ghost"],
                unselected_hover_color=CLR["ghost_h"], text_color=CLR["text"],
            )
            self.seg_theme.set("深色")
            self.seg_theme.grid(row=0, column=0, padx=(0, 12))
            self.lbl_status = ctk.CTkLabel(right, text="●  就绪", font=self.f_status, width=104,
                                           height=30, corner_radius=15, fg_color=CLR["ghost"],
                                           text_color=CLR["sub"])
            self.lbl_status.grid(row=0, column=1)

        def _build_settings(self):
            card = ctk.CTkFrame(self, fg_color=CLR["card"], corner_radius=14,
                                border_width=1, border_color=CLR["line"])
            card.grid(row=1, column=0, sticky="ew", padx=24, pady=(0, 12))
            card.grid_columnconfigure(4, weight=1)

            ctk.CTkLabel(card, text="抓取日期", font=self.f_sec,
                         text_color=CLR["text"]).grid(row=0, column=0, sticky="w",
                                                      padx=(20, 12), pady=(18, 18))
            self.entry_day = ctk.CTkEntry(card, width=150, height=38, font=self.f_entry,
                                          corner_radius=9, justify="center",
                                          fg_color=CLR["log_bg"], border_color=CLR["line"],
                                          text_color=CLR["text"])
            self.entry_day.grid(row=0, column=1, padx=(0, 14), pady=(18, 18))
            self.entry_day.insert(0, self._day_default)

            quick = ctk.CTkFrame(card, fg_color="transparent")
            quick.grid(row=0, column=2, sticky="w", pady=(18, 18))
            for i, (text, off) in enumerate([("昨天", 1), ("今天", 0), ("前天", 2)]):
                ctk.CTkButton(quick, text=text, width=66, height=32, corner_radius=8,
                              font=self.f_body, fg_color=CLR["ghost"], hover_color=CLR["ghost_h"],
                              text_color=CLR["text"],
                              command=lambda o=off: self._set_quick_day(o)).grid(
                    row=0, column=i, padx=(0, 8))

            ctk.CTkLabel(card, text="格式 YYYYMMDD，默认抓取昨天", font=self.f_sub,
                         text_color=CLR["sub"]).grid(row=0, column=4, sticky="e",
                                                     padx=(0, 20), pady=(18, 18))

        def _build_actions(self):
            act = ctk.CTkFrame(self, fg_color="transparent")
            act.grid(row=2, column=0, sticky="ew", padx=24, pady=(0, 12))
            self.btn_start = ctk.CTkButton(
                act, text="开始抓取并处理", width=208, height=46, corner_radius=11,
                font=self.f_btn, fg_color=CLR["accent"], hover_color=CLR["accent_h"],
                text_color="#FFFFFF", command=self._on_start,
            )
            self.btn_start.grid(row=0, column=0)
            ctk.CTkButton(act, text="打开输出目录", width=140, height=46, corner_radius=11,
                          font=self.f_body, fg_color=CLR["ghost"], hover_color=CLR["ghost_h"],
                          text_color=CLR["text"],
                          command=self._open_output_folder).grid(row=0, column=1, padx=(12, 0))
            self.btn_check = ctk.CTkButton(
                act, text="检查更新", width=118, height=46, corner_radius=11,
                font=self.f_body, fg_color=CLR["ghost"], hover_color=CLR["ghost_h"],
                text_color=CLR["text"],
                command=lambda: self._check_update(manual=True),
            )
            self.btn_check.grid(row=0, column=2, padx=(12, 0))

        def _build_log(self):
            box = ctk.CTkFrame(self, fg_color=CLR["card"], corner_radius=14,
                               border_width=1, border_color=CLR["line"])
            box.grid(row=3, column=0, sticky="nsew", padx=24, pady=(0, 12))
            box.grid_columnconfigure(0, weight=1)
            box.grid_rowconfigure(1, weight=1)

            head = ctk.CTkFrame(box, fg_color="transparent")
            head.grid(row=0, column=0, sticky="ew", padx=18, pady=(14, 8))
            head.grid_columnconfigure(0, weight=1)
            ctk.CTkLabel(head, text="运行日志", font=self.f_sec,
                         text_color=CLR["text"]).grid(row=0, column=0, sticky="w")
            ctk.CTkButton(head, text="清空", width=58, height=26, corner_radius=7, font=self.f_body,
                          fg_color="transparent", hover_color=CLR["ghost"], text_color=CLR["sub"],
                          command=self._clear_log).grid(row=0, column=1, sticky="e")

            self.txt = ctk.CTkTextbox(box, corner_radius=10, font=self.f_log, wrap="word",
                                      fg_color=CLR["log_bg"], border_width=1,
                                      border_color=CLR["line"], text_color=CLR["text"])
            self.txt.grid(row=1, column=0, sticky="nsew", padx=18, pady=(0, 16))
            self.txt.configure(state="disabled")

        def _build_footer(self):
            foot = ctk.CTkFrame(self, fg_color="transparent")
            foot.grid(row=4, column=0, sticky="ew", padx=26, pady=(0, 16))
            foot.grid_columnconfigure(1, weight=1)
            self.lbl_foot = ctk.CTkLabel(foot, text="就绪，等待开始", font=self.f_body,
                                         text_color=CLR["sub"])
            self.lbl_foot.grid(row=0, column=0, sticky="w")
            self.progress = ctk.CTkProgressBar(foot, width=240, height=6, corner_radius=3,
                                               progress_color=RUN_CLR, fg_color=CLR["ghost"])
            self.progress.set(0)
            self.progress.grid(row=0, column=2, sticky="e")
            ctk.CTkLabel(foot, text=f"v{current_app_version()}", font=self.f_sub,
                         text_color=CLR["sub"]).grid(row=0, column=3, sticky="e", padx=(14, 0))

        # ---------------- 交互 ----------------
        def _switch_theme(self, value):
            ctk.set_appearance_mode("dark" if value == "深色" else "light")

        def _set_quick_day(self, offset):
            day = (date.today() - timedelta(days=offset)).strftime("%Y%m%d")
            self.entry_day.delete(0, "end")
            self.entry_day.insert(0, day)

        def _clear_log(self):
            self.txt.configure(state="normal")
            self.txt.delete("1.0", "end")
            self.txt.configure(state="disabled")

        def _append_log(self, text):
            self.txt.configure(state="normal")
            self.txt.insert("end", text + "\n")
            self.txt.see("end")
            self.txt.configure(state="disabled")

        def _drain(self):
            try:
                while True:
                    kind, payload = self._q.get_nowait()
                    if kind == "log":
                        self._append_log(payload)
                    elif kind == "done":
                        self._on_finished(payload)
                    elif kind == "error":
                        self._on_failed(payload[0], payload[1])
                    elif kind == "ask_login":
                        self._show_login(payload)
                    elif kind == "update":
                        self._show_update_dialog(payload)
                    elif kind == "update_progress":
                        if self._update_dialog is not None:
                            self._update_dialog._on_progress(payload[0], payload[1])
                    elif kind == "update_ready":
                        if self._update_dialog is not None:
                            self._update_dialog._on_ready(payload)
                    elif kind == "update_error":
                        if self._update_dialog is not None:
                            self._update_dialog._on_error(payload)
                    elif kind == "check_state":
                        self.btn_check.configure(state=payload[0], text=payload[1])
            except queue.Empty:
                pass
            self.after(80, self._drain)

        def _ask_login_threadsafe(self, tip="", error="", default_account=""):
            """抓取线程请求登录：投递到主线程弹框并等待结果（Tk 只能在主线程操作）"""
            req = {"tip": tip, "error": error, "account": default_account,
                   "res": None, "event": threading.Event()}
            self._q.put(("ask_login", req))
            req["event"].wait()
            return req["res"]

        def _show_login(self, req):
            """在主线程弹出登录对话框（由 _drain 调度）"""
            try:
                req["res"] = LoginDialog(
                    self, tip=req.get("tip", ""), error=req.get("error", ""),
                    default_account=req.get("account", ""),
                ).wait()
            except Exception as e:
                log(f"[警告] 登录窗口打开失败: {e}")
                req["res"] = None
            finally:
                req["event"].set()

        def _open_output_folder(self):
            try:
                path = open_folder(os.path.join(script_dir(), "replay核查"))
                self._append_log(f"{datetime.now():%H:%M:%S}  [提示] 已在资源管理器中打开输出目录: {path}")
            except Exception as e:
                self._append_log(f"{datetime.now():%H:%M:%S}  [警告] 打开目录失败: {e}")

        def _preheat(self):
            """后台预热 pandas，降低首次点击后的等待"""
            try:
                import pandas  # noqa: F401
            except Exception:
                pass

        # ---------------- 在线更新 ----------------
        def _check_update(self, manual=False):
            """检测新版本；manual=True 表示手动点击（会输出检测过程与结果）"""
            try:
                self.btn_check.configure(state="disabled", text="检测中…")
            except Exception:
                pass
            if manual:
                self._append_log(f"{datetime.now():%H:%M:%S}  [更新] 正在检查新版本…")
            threading.Thread(target=self._update_worker, args=(manual,), daemon=True).start()

        def _update_worker(self, manual):
            try:
                status, payload = check_for_update()
            except Exception as e:      # check_for_update 内部已兜底，这里再防一手
                status, payload = "error", f"{type(e).__name__}: {e}"
            ts = f"{datetime.now():%H:%M:%S}"
            if status == "update":
                self._q.put(("update", payload))
            elif manual:
                if status == "latest":
                    self._q.put(("log", f"{ts}  [更新] 已是最新版本 v{payload.get('version', '')}"))
                else:
                    self._q.put(("log", f"{ts}  [更新] 检测失败：{payload}"))
            self._q.put(("check_state", ("normal", "检查更新")))

        def _show_update_dialog(self, info):
            self._append_log(
                f"{datetime.now():%H:%M:%S}  [更新] 发现新版本 v{info['version']}"
                f"（当前 v{current_app_version()}），可点击「立即更新」自动升级")
            try:
                if self._update_dialog is not None:
                    try:
                        self._update_dialog.destroy()
                    except Exception:
                        pass
                self._update_dialog = UpdateDialog(self, info)
            except Exception as e:
                log(f"[警告] 更新窗口打开失败: {e}")

        def _quit_for_update(self):
            """在线更新完成：退出当前进程，由新版本接管"""
            try:
                self.destroy()
            except Exception:
                pass
            _exit_now(0)

        # ---------------- 运行 ----------------
        def _on_start(self):
            if self._running:
                return
            day = re.sub(r"[^0-9]", "", self.entry_day.get())
            try:
                datetime.strptime(day, "%Y%m%d")
            except ValueError:
                self.entry_day.configure(border_color=ERR_CLR)
                self.after(1500, lambda: self.entry_day.configure(border_color=CLR["line"]))
                self._append_log(f"{datetime.now():%H:%M:%S}  [提示] 日期格式不正确，请按 20260924 格式输入")
                self._set_state("日期有误", ERR_CLR, "请检查抓取日期格式")
                return

            self._running = True
            self.btn_start.configure(state="disabled", text="处理中…")
            self.progress.configure(mode="indeterminate")
            self.progress.start()
            self._set_state("运行中", RUN_CLR, "正在抓取并处理，请稍候")
            self._append_log("")
            threading.Thread(target=self._worker, args=(day,), daemon=True).start()

        def _worker(self, day):
            try:
                res = run_pipeline(day)
            except PipelineError as e:
                self._q.put(("error", (e.stage, str(e))))
                return
            except Exception as e:
                report_exception(e, stage="unknown", day=str(day))
                self._q.put(("error", ("unknown", str(e))))
                return
            self._q.put(("done", res))

        def _set_state(self, text, color, footer=None):
            self.lbl_status.configure(text=f"●  {text}", text_color=color)
            if footer:
                self.lbl_foot.configure(text=footer)

        def _reset_progress(self, done=True):
            try:
                self.progress.stop()
            except Exception:
                pass
            self.progress.configure(mode="determinate")
            self.progress.set(1 if done else 0)

        def _on_finished(self, res):
            self._running = False
            self.btn_start.configure(state="normal", text="开始抓取并处理")
            if res.get("status") == "ok":
                self._reset_progress(True)
                stat = res.get("stat", {})
                self._set_state("已完成", OK_CLR,
                                f"完成：{stat.get('rows_after_dedup', 0)} 行，输出 {os.path.basename(stat.get('out_path', ''))}")
            else:
                self._reset_progress(False)
                self._set_state("无数据", WARN_CLR, "该日期没有已核查记录")
            try:
                ResultWindow(self, res)
            except Exception as e:
                log(f"[警告] 结果窗口打开失败: {e}")

        def _on_failed(self, stage, msg):
            self._running = False
            self.btn_start.configure(state="normal", text="开始抓取并处理")
            self._reset_progress(False)
            stage_text = {"fetch": "抓取失败", "csv": "导出失败", "process": "处理失败"}.get(stage, "运行失败")
            self._set_state(stage_text, ERR_CLR, f"{stage_text}，详见日志")
            try:
                import tkinter.messagebox as mb
                mb.showerror("运行失败", f"{stage_text}\n\n{msg}")
            except Exception:
                pass

        def _on_close(self):
            if self._running:
                try:
                    import tkinter.messagebox as mb
                    if not mb.askyesno("确认退出", "任务正在运行中，确定要退出吗？"):
                        return
                except Exception:
                    pass
            self.destroy()

    def launch_gui(selfcheck=False):
        """启动图形界面；selfcheck=True 时短暂启动后自动退出（用于打包自检）"""
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")
        app = ReplayApp(selfcheck=selfcheck)
        if selfcheck:
            app.after(1800, lambda: (app.destroy(), _console("SELFCHECK-OK")))
        app.mainloop()
        return 0


# ==================================================================
#                              入口
# ==================================================================
def _parse_args(argv):
    """返回 (day, selfcheck, action)；
    day 为 YYYYMMDD 或 None，action ∈ {None, "check_update", "version"}"""
    day, selfcheck, action, rest = None, False, None, []
    i = 0
    while i < len(argv):
        a = argv[i]
        if a in ("--selfcheck", "-s"):
            selfcheck = True
        elif a in ("--check-update",):
            action = "check_update"
        elif a in ("--version", "-V"):
            action = "version"
        elif a in ("--cli", "--date", "-d"):
            i += 1
            if i < len(argv):
                day = re.sub(r"[^0-9]", "", argv[i])
        elif a.startswith("--date="):
            day = re.sub(r"[^0-9]", "", a.split("=", 1)[1])
        elif a in ("--gui", "-g"):
            pass
        else:
            rest.append(a)
        i += 1
    if day is None and rest:
        day = re.sub(r"[^0-9]", "", rest[0])
    return day, selfcheck, action


def main():
    # 自检开关: 只执行一次「打开输出目录」动作后退出（用于打包产物离线验证，不影响正常使用）
    if "--open-out" in sys.argv[1:]:
        try:
            open_folder(os.path.join(script_dir(), "replay核查"))
            return 0
        except Exception:
            return 1

    day_arg, selfcheck, action = _parse_args(sys.argv[1:])

    # --version：打印当前版本号后退出（便于 CI 校验注入结果）
    if action == "version":
        _console(current_app_version())
        return 0

    # --check-update：只检测更新并打印结果，不启动界面
    if action == "check_update":
        status, payload = check_for_update()
        if status == "update":
            _console(f"NEW={payload['version']}  FILE={payload['asset_name']}")
        elif status == "latest":
            _console(f"LATEST  current={current_app_version()}")
        else:
            _console(f"ERROR  {payload}")
        return 0 if status in ("update", "latest") else 1

    # 上次在线更新的收尾：清理旧版本文件（后台进行，不影响启动）
    if not selfcheck:
        threading.Thread(target=cleanup_after_update, daemon=True).start()

    # 1. 命令行静默模式（传日期）
    if day_arg:
        try:
            datetime.strptime(day_arg, "%Y%m%d")
        except ValueError:
            _console(f"日期参数格式错误: {day_arg}，请使用 20260924 格式。")
            return 2
        try:
            res = run_pipeline(day_arg)
        except PipelineError:
            return 1
        except Exception as e:
            report_exception(e, stage="startup")
            _console(f"[错误] {e}")
            return 1
        return 0 if res.get("status") in ("ok", "empty") else 1

    # 2. 图形界面模式（有 customtkinter）
    if ctk is not None:
        try:
            return launch_gui(selfcheck=selfcheck)
        except Exception as e:
            report_exception(e, stage="gui")
            try:
                with open(os.path.join(script_dir(), "error.log"), "a", encoding="utf-8") as f:
                    f.write(f"{datetime.now():%Y-%m-%d %H:%M:%S} 界面异常: {e}\n")
            except Exception:
                pass
            return 1

    # 3. 降级：原生弹窗交互
    default_day = (date.today() - timedelta(days=1)).strftime("%Y%m%d")
    day = ask_date(default_day)
    if not day:
        log("已取消。")
        return 1
    try:
        res = run_pipeline(day)
    except PipelineError as e:
        popup("运行失败", str(e), kind="error")
        return 1
    if res.get("status") == "empty":
        popup("无数据", f"{res['start_at']} ~ {res['end_at']} 没有「已核查」记录。")
        return 0
    stat = res["stat"]
    popup(
        "文档处理完成",
        f"CSV 记录: {stat['rows_before']} 条\n"
        f"筛选: {stat['rows_before']} → {stat['rows_after_filter']} 行\n"
        f"去重: {stat['rows_before_dedup']} → {stat['rows_after_dedup']} 行（按「{DEDUP_COLUMN_KEYWORD}」）\n"
        f"排序: 按「分数」降序\n"
        f"─────────────────\n"
        f"输出文件:\n{stat['out_path']}",
    )
    return 0


if __name__ == "__main__":
    _rc = 0
    try:
        _rc = main()
    except SystemExit as _e:
        _rc = _e.code
    except Exception as e:
        report_exception(e, stage="startup")
        _console(f"[错误] {e}")
        _rc = 1
    # 收尾即退出：关窗口 / 任务跑完后进程立刻消失
    _exit_now(_rc)
